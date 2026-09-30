import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIG
# ============================================================

INPUT_FILE = Path("data/validated_kpis.json")
OUTPUT_FILE = Path("data/TrustRAG_Benchmarking_Report.xlsx")


# ============================================================
# LOAD DATA
# ============================================================

with open(INPUT_FILE, "r", encoding="utf-8") as f:
    companies = json.load(f)


# ============================================================
# HELPERS
# ============================================================

def metric_value(company, metric):
    obj = company.get(metric, {})
    return obj.get("value")


def metric_status(company, metric):
    obj = company.get(metric, {})
    return obj.get("status")


def metric_source(company, metric):
    obj = company.get(metric, {})
    page = obj.get("source_page")
    evidence = obj.get("evidence")

    if page and evidence:
        return f"p.{page}: {evidence}"

    return "Not available"


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


def autofit(ws):
    for column in ws.columns:

        max_length = 0
        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 45)


# ============================================================
# CREATE WORKBOOK
# ============================================================

wb = Workbook()

# Remove default sheet
default_sheet = wb.active
wb.remove(default_sheet)


# ============================================================
# SHEET 1 — EXECUTIVE SUMMARY
# ============================================================

ws = wb.create_sheet("Executive Summary")

ws.append([
    "TrustRAG — Corporate Benchmarking Report"
])

ws["A1"].font = Font(
    bold=True,
    size=16
)

ws.append([])

ws.append([
    "Company",
    "Revenue (₹ Cr)",
    "Revenue Growth",
    "Operating Margin",
    "PAT (₹ Cr)",
    "PAT Margin",
    "ROE",
    "Employees",
    "Revenue / Employee"
])

style_header(ws, 3)

for company in companies:

    ws.append([
        company["company"],
        metric_value(company, "revenue"),
        metric_value(company, "revenue_growth"),
        metric_value(company, "operating_margin"),
        metric_value(company, "pat"),
        metric_value(company, "pat_margin"),
        metric_value(company, "roe"),
        metric_value(company, "employees"),
        metric_value(company, "revenue_per_employee"),
    ])


# ============================================================
# SHEET 2 — FINANCIAL BENCHMARK
# ============================================================

ws = wb.create_sheet("Financial Benchmark")

ws.append([
    "Company",
    "Revenue (₹ Cr)",
    "Revenue Growth (%)",
    "Operating Margin (%)",
    "PAT (₹ Cr)",
    "PAT Margin (%)",
    "EPS (₹)",
    "ROE (%)",
    "Cash / Investments (₹ Cr)",
    "FCF (₹ Cr)",
    "Capex (₹ Cr)"
])

style_header(ws)

for company in companies:

    ws.append([
        company["company"],
        metric_value(company, "revenue"),
        metric_value(company, "revenue_growth"),
        metric_value(company, "operating_margin"),
        metric_value(company, "pat"),
        metric_value(company, "pat_margin"),
        metric_value(company, "eps"),
        metric_value(company, "roe"),
        metric_value(company, "cash_liquid_assets"),
        metric_value(company, "free_cash_flow"),
        metric_value(company, "capex"),
    ])


# ============================================================
# SHEET 3 — WORKFORCE
# ============================================================

ws = wb.create_sheet("Workforce")

ws.append([
    "Company",
    "Employees",
    "Employee Growth (%)",
    "Attrition (%)",
    "Revenue / Employee",
    "Attrition Definition"
])

style_header(ws)

for company in companies:

    attrition = company.get(
        "attrition",
        {}
    )

    ws.append([
        company["company"],
        metric_value(company, "employees"),
        metric_value(company, "employee_growth"),
        metric_value(company, "attrition"),
        metric_value(company, "revenue_per_employee"),
        attrition.get("notes"),
    ])


# ============================================================
# SHEET 4 — KPI STATUS
# ============================================================

ws = wb.create_sheet("KPI Status")

metrics = [
    "revenue",
    "revenue_growth",
    "operating_margin",
    "pat",
    "eps",
    "roe",
    "free_cash_flow",
    "capex",
    "cash_liquid_assets",
    "employees",
    "attrition",
]

ws.append([
    "Company"
] + metrics)

style_header(ws)

for company in companies:

    row = [company["company"]]

    for metric in metrics:

        row.append(
            metric_status(
                company,
                metric
            )
        )

    ws.append(row)


# ============================================================
# SHEET 5 — EVIDENCE
# ============================================================

ws = wb.create_sheet("Evidence")

ws.append([
    "Company",
    "KPI",
    "Value",
    "Unit",
    "Fiscal Year",
    "Source Page",
    "Evidence",
    "Notes"
])

style_header(ws)

for company in companies:

    for metric in metrics:

        obj = company.get(
            metric,
            {}
        )

        ws.append([
            company["company"],
            metric,
            obj.get("value"),
            obj.get("unit"),
            obj.get("fiscal_year"),
            obj.get("source_page"),
            obj.get("evidence"),
            obj.get("notes"),
        ])


# ============================================================
# SHEET 6 — COMPARABILITY
# ============================================================

ws = wb.create_sheet("Comparability")

ws.append([
    "KPI",
    "Comparability",
    "Reason"
])

style_header(ws)

comparability_metrics = set()

for company in companies:

    for metric, details in company.get(
        "comparability",
        {}
    ).items():

        comparability_metrics.add(metric)

# Use the first available definition
for metric in sorted(comparability_metrics):

    details = None

    for company in companies:

        if metric in company.get(
            "comparability",
            {}
        ):

            details = company[
                "comparability"
            ][metric]

            break

    ws.append([
        metric,
        details.get("level"),
        details.get("reason"),
    ])


# ============================================================
# FORMATTING
# ============================================================

for ws in wb.worksheets:

    ws.freeze_panes = "A2"

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    autofit(ws)


# ============================================================
# SAVE
# ============================================================

OUTPUT_FILE.parent.mkdir(
    parents=True,
    exist_ok=True
)

wb.save(OUTPUT_FILE)

print()
print("=" * 70)
print("TRUSTRAG BENCHMARKING REPORT")
print("=" * 70)

print(
    f"Companies: {len(companies)}"
)

print(
    f"Sheets created: {len(wb.sheetnames)}"
)

print()

for sheet in wb.sheetnames:
    print(f"  ✓ {sheet}")

print()
print(
    f"Saved to: {OUTPUT_FILE}"
)